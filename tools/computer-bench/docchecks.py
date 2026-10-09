#!/usr/bin/env python3
# Preparación y comprobaciones de las tareas doc-* del banco computer-bench (Python 3.9+).
# Uso: docchecks.py setup <tarea> <escritorio> | check <tarea> <escritorio>
# `check` imprime una línea JSON {"pass": bool, "why": str}.
import hashlib
import json
import os
import re
import sys


def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        h.update(f.read())
    return h.hexdigest()


def res(ok, why):
    print(json.dumps({"pass": bool(ok), "why": why}, ensure_ascii=False))
    sys.exit(0)


def nums_in(text):
    return [int(x.replace(".", "")) for x in re.findall(r"\d[\d.]*", text) if x.replace(".", "").isdigit()]


# ---------- PDF mínimo sin dependencias (para el setup) ----------
def make_pdf(path, pages_text):
    objs = []
    n = len(pages_text)
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join("%d 0 R" % (4 + 2 * i) for i in range(n))
    objs.append(("<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, n)).encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, t in enumerate(pages_text):
        objs.append(("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (5 + 2 * i)).encode())
        stream = ("BT /F1 24 Tf 72 700 Td (%s) Tj ET" % t).encode()
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = b"%PDF-1.4\n"
    offs = []
    for i, o in enumerate(objs):
        offs.append(len(out))
        out += ("%d 0 obj\n" % (i + 1)).encode() + o + b"\nendobj\n"
    x = len(out)
    out += ("xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)).encode()
    for o in offs:
        out += ("%010d 00000 n \n" % o).encode()
    out += ("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, x)).encode()
    with open(path, "wb") as f:
        f.write(out)


# ---------- setups ----------
def setup_word_editar(d):
    import docx
    doc = docx.Document()
    doc.add_heading("Contrato de servicios", 0)
    doc.add_paragraph("Entre la empresa Carsav SpA y el cliente Juan Pérez se acuerda lo siguiente.")
    doc.add_paragraph("Cláusula 1: Juan Pérez recibirá el servicio de mantención mensual.")
    doc.add_paragraph("Cláusula 2: El pago será de 450000 pesos mensuales.")
    p = os.path.join(d, "bench-contrato.docx")
    doc.save(p)
    with open(os.path.join(d, ".trabajo-bench-contrato.sha256"), "w") as f:
        f.write(sha(p))


def setup_excel_editar(d):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Precios"
    ws.append(["Producto", "Cantidad", "Precio"])
    for row in [("Cuaderno", 10, 1500), ("Lápiz", 25, 300), ("Mochila", 4, 18000)]:
        ws.append(list(row))
    wb.save(os.path.join(d, "bench-precios.xlsx"))


def setup_ppt_editar(d):
    from pptx import Presentation
    prs = Presentation()
    for t, b in [("Equipo Alfa", "Quiénes somos"), ("Objetivos 2026", "Crecer un 20 por ciento"), ("Resultados", "Ventas al alza")]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
        s.placeholders[1].text = b
    prs.save(os.path.join(d, "bench-equipo.pptx"))


def setup_pdf_unir(d):
    make_pdf(os.path.join(d, "bench-parte-a.pdf"), ["PAGINA-A1", "PAGINA-A2"])
    make_pdf(os.path.join(d, "bench-parte-b.pdf"), ["PAGINA-B1", "PAGINA-B2", "PAGINA-B3"])


# ---------- checks ----------
def need(path):
    if not os.path.exists(path):
        res(False, "no existe " + os.path.basename(path))


def check_word_informe(d):
    import docx
    p = os.path.join(d, "bench-informe-ventas.docx")
    need(p)
    doc = docx.Document(p)
    if not doc.tables:
        res(False, "el documento no tiene tabla")
    t = max(doc.tables, key=lambda x: len(x.rows))
    if len(t.rows) != 5:
        res(False, "la tabla tiene %d filas (se esperaban 5: encabezado, 3 meses y total)" % len(t.rows))
    last = " ".join(c.text for c in t.rows[-1].cells)
    if not re.search(r"total", last, re.I) or 560 not in nums_in(last):
        res(False, "la última fila no es el total 560: «%s»" % last.strip())
    body = [x for x in doc.paragraphs if x.text.strip()]
    heads = [x for x in body if x.style.name.lower().startswith(("title", "heading", "título", "titulo"))]
    if not heads:
        res(False, "no hay título (estilo Title/Heading)")
    normal = [x for x in body if x not in heads]
    if len(normal) < 3:
        res(False, "el resumen tiene menos de 3 párrafos/líneas (%d)" % len(normal))
    res(True, "título, resumen (%d líneas) y tabla de 5 filas con total 560" % len(normal))


def check_word_editar(d):
    import docx
    orig = os.path.join(d, "bench-contrato.docx")
    shaf = os.path.join(d, ".trabajo-bench-contrato.sha256")
    need(orig)
    if not os.path.exists(shaf):
        res(False, "falta el sha256 guardado por el setup (ejecuta setup)")
    if sha(orig) != open(shaf).read().strip():
        res(False, "el ORIGINAL bench-contrato.docx fue modificado")
    cands = [f for f in os.listdir(d) if f.startswith("bench-contrato") and f.endswith(".docx") and f != "bench-contrato.docx"]
    if not cands:
        res(False, "no hay copia nueva bench-contrato*.docx")
    why = []
    for f in cands:
        doc = docx.Document(os.path.join(d, f))
        text = "\n".join(p.text for p in doc.paragraphs)
        if "Juan Pérez" in text:
            why.append(f + ": aún contiene «Juan Pérez»")
        elif "María González" not in text:
            why.append(f + ": no contiene «María González»")
        elif "450000" not in text.replace(".", "") or "Carsav" not in text:
            why.append(f + ": perdió contenido original")
        else:
            res(True, "copia %s con el nombre reemplazado y el original intacto" % f)
    res(False, "; ".join(why))


def check_excel_ventas(d):
    import openpyxl
    p = os.path.join(d, "bench-ventas.xlsx")
    need(p)
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    vals = {}
    for row in ws.iter_rows():
        for c in row:
            vals[c.coordinate] = c.value
    sums = [(k, v) for k, v in vals.items() if isinstance(v, str) and re.match(r"=\s*SUM\(", v, re.I)]
    if not sums:
        res(False, "no hay ninguna fórmula SUMA/SUM (valores escritos a mano?)")
    good = False
    for k, f in sums:
        m = re.search(r"SUM\(\s*\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)\s*\)", f, re.I)
        if not m:
            continue
        rng = ws["%s%s:%s%s" % (m.group(1), m.group(2), m.group(3), m.group(4))]
        total = 0
        for r in rng:
            for c in r:
                if isinstance(c.value, (int, float)):
                    total += c.value
        if total == 560:
            good = True
    if not good:
        res(False, "ninguna SUMA cubre los datos y suma 560: %s" % [f for _, f in sums])
    n_charts = sum(len(w._charts) for w in wb.worksheets)
    if n_charts < 1:
        res(False, "no hay gráfico")
    res(True, "SUMA correcta (560) y %d gráfico(s)" % n_charts)


def check_excel_editar(d):
    import openpyxl
    p = os.path.join(d, "bench-precios.xlsx")
    need(p)
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    orig = [("Cuaderno", 10, 1500), ("Lápiz", 25, 300), ("Mochila", 4, 18000)]
    for i, row in enumerate(orig, start=2):
        got = (ws.cell(i, 1).value, ws.cell(i, 2).value, ws.cell(i, 3).value)
        if got != row:
            res(False, "datos originales alterados en fila %d: %s" % (i, got))
    h = ws.cell(1, 4).value
    if not h or "total" not in str(h).lower():
        res(False, "la columna D no tiene encabezado «Total» (tiene %r)" % h)
    for i in (2, 3, 4):
        f = ws.cell(i, 4).value
        if not (isinstance(f, str) and f.startswith("=") and re.search(r"B%d" % i, f) and re.search(r"C%d" % i, f)):
            res(False, "D%d no es una fórmula con B%d y C%d: %r" % (i, i, i, f))
    res(True, "columna Total con fórmulas =B*C en las 3 filas y datos originales intactos")


def check_ppt_nueva(d):
    from pptx import Presentation
    p = os.path.join(d, "bench-presentacion.pptx")
    need(p)
    prs = Presentation(p)
    if len(prs.slides) != 4:
        res(False, "tiene %d diapositivas (se esperaban 4)" % len(prs.slides))
    charts = sum(1 for s in prs.slides for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart)
    if charts < 1:
        res(False, "ninguna diapositiva tiene un gráfico real")
    for i, s in enumerate(prs.slides, 1):
        txt = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame).strip()
        if not txt:
            res(False, "la diapositiva %d no tiene texto" % i)
    res(True, "4 diapositivas con texto y %d gráfico(s)" % charts)


def check_ppt_editar(d):
    from pptx import Presentation
    p = os.path.join(d, "bench-equipo.pptx")
    need(p)
    prs = Presentation(p)
    if len(prs.slides) != 4:
        res(False, "tiene %d diapositivas (se esperaban 4)" % len(prs.slides))
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    for i, t in enumerate(["Equipo Alfa", "Objetivos 2026", "Resultados"]):
        if t not in texts[i]:
            res(False, "la diapositiva %d original cambió (falta «%s»)" % (i + 1, t))
    if "próximos pasos" not in texts[3].lower():
        res(False, "la diapositiva 4 no se titula «Próximos pasos»")
    lines = [ln for sh in prs.slides[3].shapes if sh.has_text_frame for ln in sh.text_frame.text.split("\n") if ln.strip()]
    if len(lines) < 4:
        res(False, "la diapositiva nueva no tiene título + 3 viñetas (%d líneas)" % len(lines))
    res(True, "4 diapositivas, originales intactas y «Próximos pasos» con viñetas")


def pdf_text(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    return r, [(pg.extract_text() or "") for pg in r.pages]


def check_pdf_tabla(d):
    p = os.path.join(d, "bench-tabla.pdf")
    need(p)
    r, texts = pdf_text(p)
    t = "\n".join(texts)
    for w in ("Producto", "Precio", "Manzanas", "Peras", "Uvas"):
        if w.lower() not in t.lower():
            res(False, "el PDF no contiene «%s»" % w)
    n = set(nums_in(t))
    if not {1200, 800, 2500} <= n:
        res(False, "faltan los precios 1200/800/2500 en el PDF (hay %s)" % sorted(n)[:12])
    res(True, "PDF de %d página(s) con la tabla Producto/Precio y los 3 productos" % len(r.pages))


def check_pdf_unir(d):
    p = os.path.join(d, "bench-unido.pdf")
    need(p)
    r, texts = pdf_text(p)
    order = [re.search(r"PAGINA-[AB]\d", t).group(0) if re.search(r"PAGINA-[AB]\d", t) else "?" for t in texts]
    want = ["PAGINA-A1", "PAGINA-A2", "PAGINA-B1", "PAGINA-B3"]
    if order != want:
        res(False, "páginas %s, se esperaba %s" % (order, want))
    res(True, "4 páginas en el orden correcto, sin la página 4 original (B2)")


def main():
    if len(sys.argv) != 4:
        print("uso: docchecks.py setup|check <tarea> <escritorio>")
        sys.exit(2)
    mode, task, d = sys.argv[1:4]
    fn = globals().get(("setup_" if mode == "setup" else "check_") + task.replace("-", "_"))
    if fn is None:
        res(False, "tarea desconocida: " + task)
    try:
        fn(d)
    except SystemExit:
        raise
    except Exception as e:  # archivo corrupto o librería ausente
        res(False, "error al leer: %s: %s" % (type(e).__name__, e))
    if mode == "setup":
        print("ok")


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""Pruebas del prefijo _xlfn. en fórmulas modernas. Ejecutar: python -m unittest tests/test_xlfn.py"""
import json
import os
import subprocess
import sys
import tempfile
import unittest
import warnings
import zipfile

warnings.filterwarnings("ignore")
from openpyxl import load_workbook

SCRIPTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts")
sys.path.insert(0, SCRIPTS)
from xlsx_common import add_xlfn_prefix  # noqa: E402


def run(script, *args):
    return subprocess.run([sys.executable, os.path.join(SCRIPTS, script)] + list(args), capture_output=True, text=True)


class AddXlfnPrefix(unittest.TestCase):
    def test_funciones_modernas(self):
        casos = {
            '=TEXTJOIN(",",TRUE,A1:A3)': '=_xlfn.TEXTJOIN(",",TRUE,A1:A3)',
            "=XLOOKUP(A1,B:B,C:C)": "=_xlfn.XLOOKUP(A1,B:B,C:C)",
            "=IFS(A1>1,1,TRUE,0)": "=_xlfn.IFS(A1>1,1,TRUE,0)",
            "=DAYS(B1,A1)": "=_xlfn.DAYS(B1,A1)",
            "=XOR(A1,B1)": "=_xlfn.XOR(A1,B1)",
            "=STDEV.S(A1:A9)": "=_xlfn.STDEV.S(A1:A9)",
            "=CONCAT(A1,B1)": "=_xlfn.CONCAT(A1,B1)",
            "=MAXIFS(A:A,B:B,1)+MINIFS(A:A,B:B,1)": "=_xlfn.MAXIFS(A:A,B:B,1)+_xlfn.MINIFS(A:A,B:B,1)",
            "=SWITCH(A1,1,2)": "=_xlfn.SWITCH(A1,1,2)",
            "=XMATCH(1,A:A)": "=_xlfn.XMATCH(1,A:A)",
            "=textjoin(\",\",TRUE,A1)": "=_xlfn.TEXTJOIN(\",\",TRUE,A1)",
            "=IF(A1>0,TEXTJOIN(\",\",1,B1),\"\")": "=IF(A1>0,_xlfn.TEXTJOIN(\",\",1,B1),\"\")",
            "=TEXTJOIN (\",\",1,B1)": "=_xlfn.TEXTJOIN (\",\",1,B1)",
        }
        for entrada, esperado in casos.items():
            self.assertEqual(add_xlfn_prefix(entrada), esperado, entrada)

    def test_no_toca_lo_que_no_debe(self):
        intactos = [
            "=_xlfn.TEXTJOIN(\",\",TRUE,A1)",
            "=_xlfn.STDEV.S(A1:A9)",
            "=SUM(A1:A3)",
            "=DAYS360(A1,B1)",
            "=STDEV(A1:A9)",
            '=SUM(A1)&" usa XLOOKUP( y IFS( en texto"',
            '="a ""TEXTJOIN(""" & A1',
            "='XLOOKUP(2)'!A1",
            "=MIYXLOOKUP(A1)",
            "=IFS",
            "TEXTJOIN(A1)",
            "texto cualquiera",
            "=",
            None,
            12,
        ]
        for v in intactos:
            self.assertEqual(add_xlfn_prefix(v), v, v)

    def test_idempotente(self):
        f = "=XLOOKUP(A1,B:B,IFS(C1>1,1,TRUE,2))"
        una = add_xlfn_prefix(f)
        self.assertEqual(add_xlfn_prefix(una), una)


class RecalcCheckConoceLasModernas(unittest.TestCase):
    def test_sintaxis_acepta_funciones_con_prefijo(self):
        import recalc_check
        for f in ("=_xlfn.STDEV.S(A1:A3)", "=_xlfn.TEXTJOIN(\",\",TRUE,A1:A3)", "=STDEV.S(A1:A3)"):
            self.assertEqual(recalc_check.syntax_issues(f, None), [], f)
        self.assertTrue(recalc_check.syntax_issues("=FOO.BAR(A1)", None))


class EscrituraEnArchivo(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def formulas(self, path):
        with zipfile.ZipFile(path) as z:
            xml = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
        return xml

    def test_build_xlsx_prefija_filas_totales_y_resumen(self):
        spec = {"sheets": [{
            "name": "D", "columns": [
                {"header": "A"}, {"header": "N", "format": "integer", "total": '=XLOOKUP(1,B2:B3,B2:B3)'},
                {"header": "R", "formula": '=TEXTJOIN("-",TRUE,A{r},B{r})'}],
            "rows": [["x", 1], ["z", 2]], "total_row": True}],
            "summary": {"name": "Resumen", "kpis": [{"label": "S", "formula": "=STDEV.S({col:D:B})"}]}}
        sp = os.path.join(self.tmp, "spec.json")
        out = os.path.join(self.tmp, "o.xlsx")
        with open(sp, "w") as fh:
            json.dump(spec, fh)
        p = run("build_xlsx.py", sp, "-o", out, "--force")
        self.assertEqual(p.returncode, 0, p.stderr)
        xml = self.formulas(out)
        self.assertIn('<f>_xlfn.TEXTJOIN("-",TRUE,A2,B2)</f>', xml)
        self.assertIn("_xlfn.XLOOKUP(1,B2:B3,B2:B3)", xml)
        self.assertIn("_xlfn.STDEV.S(", xml)
        self.assertNotIn("<f>TEXTJOIN", xml)
        # openpyxl devuelve la fórmula tal como está en el archivo
        wb = load_workbook(out)
        self.assertEqual(wb["D"]["C2"].value, '=_xlfn.TEXTJOIN("-",TRUE,A2,B2)')

    def test_edit_xlsx_prefija_celdas_y_columnas(self):
        base = os.path.join(self.tmp, "b.xlsx")
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "H"
        ws.append(["A", "B"])
        ws.append([1, 2])
        wb.save(base)
        ops = [{"op": "set_cells", "sheet": "H", "cells": {"D1": "=IFS(A2>0,1,TRUE,0)"}},
               {"op": "add_column", "sheet": "H", "header": "X", "formula": "=DAYS(B{r},A{r})"}]
        out = os.path.join(self.tmp, "b2.xlsx")
        p = run("edit_xlsx.py", base, "--ops-json", json.dumps(ops), "-o", out)
        self.assertEqual(p.returncode, 0, p.stderr)
        xml = self.formulas(out)
        self.assertIn("_xlfn.IFS(A2&gt;0,1,TRUE,0)", xml)
        self.assertIn("_xlfn.DAYS(B2,A2)", xml)


if __name__ == "__main__":
    unittest.main(verbosity=2)

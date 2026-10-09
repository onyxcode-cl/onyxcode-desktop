# -*- coding: utf-8 -*-
"""Prueba de humo de la skill xlsx. Ejecutar: python -m unittest tests/test_xlsx.py  (o python tests/test_xlsx.py)"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import warnings

warnings.filterwarnings("ignore")
from openpyxl import load_workbook

SCRIPTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts")
sys.path.insert(0, SCRIPTS)
import recalc_check  # noqa: E402

HAS_LO = recalc_check.find_soffice() is not None

SPEC = {
    "sheets": [{
        "name": "Ventas", "title": "Ventas",
        "columns": [{"header": "Producto"}, {"header": "Cantidad", "format": "integer", "total": "sum"},
                    {"header": "Precio", "format": "clp"},
                    {"header": "Subtotal", "format": "clp", "formula": "=B{r}*C{r}", "total": "sum"},
                    {"header": "Fecha", "format": "date"}],
        "rows": [["A", 2, 1000, None, "2026-01-15"], ["B", 3, 2500, None, "2026-02-01"], ["C", 1, 400, None, "2026-03-09"]],
        "total_row": True,
        "conditional": [{"column": "D", "rule": "greater_than", "value": 2000}],
        "validations": [{"column": "B", "type": "whole", "min": 0, "max": 100}],
        "charts": [{"type": "column", "title": "Subtotal por producto", "x_title": "Producto", "y_title": "CLP",
                    "categories": "A", "series": ["D"]}],
    }],
    "summary": {"name": "Resumen", "kpis": [{"label": "Total", "formula": "={total:Ventas:D}", "format": "clp"}]},
}


def run(script, *args):
    p = subprocess.run([sys.executable, os.path.join(SCRIPTS, script)] + list(args), capture_output=True, text=True)
    return p.returncode, p.stdout, p.stderr


class TestXlsx(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="xlsxtest-")
        cls.spec = os.path.join(cls.tmp, "spec.json")
        with open(cls.spec, "w") as f:
            json.dump(SPEC, f)
        cls.book = os.path.join(cls.tmp, "libro.xlsx")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_1_build_inspect_edit_check(self):
        code, out, err = run("build_xlsx.py", self.spec, "-o", self.book)
        self.assertEqual(code, 0, err)
        wb = load_workbook(self.book)
        self.assertEqual(wb.sheetnames, ["Resumen", "Ventas"])
        ws = wb["Ventas"]
        self.assertEqual(ws["D4"].value, "=B4*C4")
        self.assertTrue(any(str(c.value).startswith("=SUM(") for r in ws.iter_rows() for c in r if c.value))
        self.assertEqual(len(ws._charts), 1)
        self.assertEqual(ws.freeze_panes, "A4")
        # no sobrescribe
        self.assertEqual(run("build_xlsx.py", self.spec, "-o", self.book)[0], 2)

        code, out, _ = run("inspect_xlsx.py", self.book, "--json")
        self.assertEqual(code, 0)
        info = json.loads(out)
        h = [x for x in info["hojas"] if x["nombre"] == "Ventas"][0]
        self.assertEqual(len(h["graficos"]), 1)
        self.assertEqual(h["graficos"][0]["titulo"], "Subtotal por producto")
        self.assertGreaterEqual(h["n_formulas"], 4)

        ops = [
            {"op": "set_cells", "sheet": "Ventas", "cells": {"B4": 10}},
            {"op": "add_column", "sheet": "Ventas", "header": "IVA", "header_row": 3, "formula": "=D{r}*0.19", "format": "clp"},
            {"op": "add_sheet", "name": "Notas", "rows": [["Nota"], ["Hecho con Lazunox"]]},
            {"op": "format", "sheet": "Ventas", "range": "A3:A3", "fill": "FFFF00"},
            {"op": "add_chart", "sheet": "Ventas", "type": "line", "title": "Cantidad", "categories": "A4:A6",
             "series": ["B3:B6"], "anchor": "J20"},
        ]
        opsf = os.path.join(self.tmp, "ops.json")
        with open(opsf, "w") as f:
            json.dump(ops, f)
        v2 = os.path.join(self.tmp, "libro-v2.xlsx")
        code, out, err = run("edit_xlsx.py", self.book, "--ops", opsf, "-o", v2)
        self.assertEqual(code, 0, err + out)
        w2 = load_workbook(v2)
        self.assertIn("Notas", w2.sheetnames)
        self.assertEqual(w2["Ventas"]["B4"].value, 10)
        self.assertEqual(len(w2["Ventas"]._charts), 2)  # conserva el original y añade uno
        self.assertEqual(w2["Ventas"].freeze_panes, "A4")  # conserva paneles
        self.assertEqual(w2["Ventas"]["C4"].number_format, load_workbook(self.book)["Ventas"]["C4"].number_format)
        self.assertEqual(w2["Ventas"]["F4"].value, "=D4*0.19")
        self.assertEqual(load_workbook(self.book)["Ventas"]["B4"].value, 2)  # original intacto
        self.assertEqual(run("edit_xlsx.py", self.book, "--ops", opsf, "-o", self.book)[0], 2)

        code, out, _ = run("check_xlsx.py", v2)
        self.assertEqual(code, 0, out)

    def test_2_recalc(self):
        if not os.path.exists(self.book):
            self.skipTest("depende de test_1")
        code, out, err = run("recalc_check.py", self.book, "--json")
        res = json.loads(out)
        if not HAS_LO:
            self.assertFalse(res["libreoffice"])
            self.assertTrue(res["avisos"])
            return
        self.assertEqual(code, 0, out)
        vals = {(v["hoja"], v["celda"]): v["valor"] for v in res["valores"]}
        self.assertEqual(vals[("Ventas", "D4")], 2000)
        self.assertEqual(vals[("Resumen", "B4")], 2000 + 7500 + 400)

    def test_3_recalc_detecta_errores(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Datos"
        ws["A1"] = 5
        ws["B1"] = 0
        ws["C1"] = "=A1/B1"
        ws["C2"] = "=SUM(E1:E5)"
        ws["C3"] = "=FOO(A1)"
        bad = os.path.join(self.tmp, "malo.xlsx")
        wb.save(bad)
        code, out, _ = run("recalc_check.py", bad, "--json")
        res = json.loads(out)
        self.assertEqual(code, 1)
        self.assertTrue(res["vacias"])
        self.assertTrue(res["sintaxis"])
        if HAS_LO:
            self.assertTrue(any(e["celda"] == "C1" and e["valor"] == "#DIV/0!" for e in res["errores"]))


if __name__ == "__main__":
    unittest.main(verbosity=2)

# .cowork/make_xlsx.py: libro Excel con openpyxl (si `python3 -c "import openpyxl"` funciona).
import sys
sys.path.insert(0, '.cowork/pylib')  # por si openpyxl se instaló con pip --target
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SALIDA = sys.argv[1] if len(sys.argv) > 1 else 'ventas.xlsx'  # usa un nombre NUEVO (ver la skill)
HOJAS = {
    'Ventas': [
        ['Mes', 'Producto', 'Unidades', 'Precio', 'Total'],
        ['Julio', 'Alfa', 10, 1500, '=C2*D2'],
        ['Agosto', 'Beta', 7, 2300, '=C3*D3'],
    ]
}

wb = Workbook()
wb.remove(wb.active)
for nombre, filas in HOJAS.items():
    ws = wb.create_sheet(nombre[:31])
    for fila in filas:
        ws.append(fila)
    for c in ws[1]:  # encabezado
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F3864')
        c.alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'A2'
    for i, col in enumerate(zip(*filas), start=1):  # ancho de columna según el contenido
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(len(str(v)) for v in col) + 3)
    for fila in ws.iter_rows(min_row=2):  # formato de miles con punto en es-CL lo pone Excel según la región
        for c in fila:
            if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith('=')):
                c.number_format = '#,##0'
wb.save(SALIDA)
print('creado', SALIDA)

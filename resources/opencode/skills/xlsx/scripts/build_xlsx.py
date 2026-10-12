#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Crea un libro Excel profesional (.xlsx) desde una especificación JSON.

Uso:
  python build_xlsx.py spec.json                 # escribe el archivo indicado en "output"
  python build_xlsx.py spec.json -o informe.xlsx
  python build_xlsx.py --example > spec.json     # imprime un JSON de ejemplo completo

Las fórmulas se guardan como fórmulas reales (openpyxl no las calcula: usa recalc_check.py).
Códigos de salida: 0 ok, 1 error de ejecución, 2 especificación inválida.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from xlsx_common import (DATE_KINDS, NUMERIC_KINDS, THEME, SpecError, add_chart, add_xlfn_prefix, autosize,  # noqa: E402
                         col_letter, fill, is_formula, number_format, parse_date,
                         style_header_cell, thin_border)

EXAMPLE = {
    "output": "ventas-2026.xlsx",
    "properties": {"title": "Ventas 2026", "author": "Lazunox"},
    "summary": {
        "name": "Resumen",
        "title": "Resumen de ventas 2026",
        "subtitle": "Cifras en pesos chilenos (CLP)",
        "kpis": [
            {"label": "Ventas totales", "formula": "={total:Ventas:C}", "format": "clp"},
            {"label": "Unidades vendidas", "formula": "={total:Ventas:B}", "format": "integer"},
            {"label": "Ticket promedio", "formula": "={total:Ventas:C}/{total:Ventas:B}", "format": "clp"},
        ],
        "charts": [
            {"type": "pie", "title": "Participación por producto", "data_sheet": "Ventas",
             "categories": "A", "series": ["C"], "anchor": "E3"}
        ],
    },
    "sheets": [
        {
            "name": "Ventas",
            "title": "Ventas por producto",
            "subtitle": "Primer semestre 2026",
            "tab_color": "1F3864",
            "columns": [
                {"header": "Producto", "width": 22},
                {"header": "Unidades", "format": "integer", "total": "sum"},
                {"header": "Ventas", "format": "clp", "total": "sum"},
                {"header": "Precio unitario", "format": "clp", "formula": "=C{r}/B{r}"},
                {"header": "Margen", "format": "percent", "total": "average"},
                {"header": "Fecha cierre", "format": "date"},
                {"header": "Estado"},
            ],
            "rows": [
                ["Notebook", 120, 96000000, None, 0.22, "2026-06-30", "Cerrado"],
                ["Monitor", 340, 61200000, None, 0.18, "2026-06-30", "Cerrado"],
                ["Teclado", 800, 14400000, None, 0.35, "2026-06-30", "En revisión"],
                ["Mouse", 950, 9500000, None, 0.4, "2026-06-30", "Pendiente"],
            ],
            "total_row": True,
            "total_label": "Total",
            "freeze": True,
            "autofilter": True,
            "zebra": True,
            "conditional": [
                {"column": "E", "rule": "less_than", "value": 0.2, "fill": "FFC7CE", "font_color": "9C0006"},
                {"column": "C", "rule": "data_bar", "color": "5B9BD5"},
            ],
            "validations": [
                {"column": "G", "type": "list", "values": ["Pendiente", "En revisión", "Cerrado"],
                 "extra_rows": 20},
                {"column": "E", "type": "decimal", "min": 0, "max": 1,
                 "message": "Ingrese un porcentaje entre 0% y 100%"},
            ],
            "charts": [
                {"type": "column", "title": "Ventas por producto", "x_title": "Producto",
                 "y_title": "Ventas (CLP)", "categories": "Producto", "series": ["Ventas"], "anchor": "I3"},
                {"type": "line", "title": "Unidades por producto", "x_title": "Producto",
                 "y_title": "Unidades", "categories": "A", "series": ["B"], "anchor": "I21"},
            ],
        }
    ],
}

HELP_SCHEMA = """Esquema (campos opcionales entre []):
{ "output": "archivo.xlsx", ["properties": {title, author}],
  ["summary": {name, title, subtitle, kpis:[{label, formula|value, format}], charts:[...]}],
  "sheets": [ {name, [title], [subtitle], [tab_color], [start_row],
      columns:[{header, [key], [width], [format], [align], [total], [formula]}],
      rows:[[...] | {header: valor}], [total_row], [total_label], [freeze], [autofilter], [zebra],
      [conditional], [validations], [charts], [column_widths], [notes]} ] }
Formatos: text integer number clp usd eur uf percent percent0 date date_long datetime month (o un código Excel).
Fórmulas: cadena que empieza con '='; {r} = nº de fila. En el resumen: {total:Hoja:COL} y {col:Hoja:COL}."""


def load_spec(path):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        raise SpecError("el archivo de especificación no existe: %s" % path)
    except json.JSONDecodeError as e:
        raise SpecError("el JSON no es válido (%s, línea %d, columna %d)" % (e.msg, e.lineno, e.colno))


class SheetInfo:
    def __init__(self, ws, headers, header_row, first, last, total_row):
        self.ws, self.headers, self.header_row = ws, headers, header_row
        self.first, self.last, self.total_row = first, last, total_row  # filas de datos


def safe_title(name):
    name = re.sub(r"[\[\]\*\?/\\:]", "-", str(name))[:31]
    return name or "Hoja"


def write_title(ws, spec, ncols):
    row = 1
    if spec.get("title"):
        c = ws.cell(row=row, column=1, value=spec["title"])
        c.font = Font(name=THEME["font"], size=16, bold=True, color=THEME["title_font"])
        ws.row_dimensions[row].height = 24
        row += 1
        if spec.get("subtitle"):
            s = ws.cell(row=row, column=1, value=spec["subtitle"])
            s.font = Font(name=THEME["font"], size=10, italic=True, color="595959")
            row += 1
        row += 1  # fila en blanco
    return row


def build_table(wb, spec):
    name = safe_title(spec.get("name", "Hoja"))
    if name in wb.sheetnames:
        raise SpecError("hay dos hojas con el nombre %r" % name)
    ws = wb.create_sheet(name)
    if spec.get("tab_color"):
        ws.sheet_properties.tabColor = spec["tab_color"]
    cols = spec.get("columns")
    if not cols:
        raise SpecError("la hoja %r no tiene 'columns'" % name)
    headers = [str(c.get("header", "")) for c in cols]
    keys = [c.get("key", c.get("header")) for c in cols]
    hrow = int(spec.get("start_row") or write_title(ws, spec, len(cols)))
    if spec.get("start_row"):
        write_title(ws, spec, len(cols))
    ws.sheet_view.showGridLines = bool(spec.get("gridlines", False))
    for i, h in enumerate(headers, 1):
        style_header_cell(ws.cell(row=hrow, column=i, value=h))
    ws.row_dimensions[hrow].height = 22
    rows = spec.get("rows") or []
    r = hrow
    zebra = spec.get("zebra", True)
    for ridx, row in enumerate(rows):
        r += 1
        if isinstance(row, dict):
            vals = [row.get(k, row.get(h)) for k, h in zip(keys, headers)]
        else:
            vals = list(row) + [None] * (len(cols) - len(row))
            if len(row) > len(cols):
                raise SpecError("la fila %d de %r tiene %d valores y hay %d columnas" % (ridx + 1, name, len(row), len(cols)))
        for ci, (col, v) in enumerate(zip(cols, vals), 1):
            if v is None and col.get("formula"):
                v = col["formula"]
            kind = str(col.get("format", "")).lower()
            if kind in DATE_KINDS:
                v = parse_date(v)
            if is_formula(v):
                v = add_xlfn_prefix(v.replace("{r}", str(r)))
            c = ws.cell(row=r, column=ci, value=v)
            c.font = Font(name=THEME["font"], size=11)
            c.border = thin_border()
            fmt = number_format(col.get("format"))
            if fmt:
                c.number_format = fmt
            align = col.get("align")
            if not align:
                is_num = kind in NUMERIC_KINDS or isinstance(v, (int, float)) and not isinstance(v, bool)
                align = "right" if is_num or kind in DATE_KINDS else "left"
            c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=bool(col.get("wrap")),
                                    indent=1 if align in ("left", "right") else 0)
            if zebra and ridx % 2 == 1:
                c.fill = fill(THEME["zebra"])
    first, last = hrow + 1, r
    total_row = None
    want_total = spec.get("total_row") or any(c.get("total") for c in cols)
    if want_total and rows:
        total_row = last + 1
        for ci, col in enumerate(cols, 1):
            L = get_column_letter(ci)
            t = col.get("total")
            c = ws.cell(row=total_row, column=ci)
            if t:
                if str(t).startswith("="):
                    c.value = add_xlfn_prefix(str(t).replace("{r}", str(total_row)))
                else:
                    fn = {"sum": "SUM", "average": "AVERAGE", "avg": "AVERAGE", "count": "COUNT",
                          "min": "MIN", "max": "MAX"}.get(str(t).lower())
                    if not fn:
                        raise SpecError("total %r no válido en %r (usa sum, average, count, min, max o una fórmula)" % (t, col.get("header")))
                    c.value = "=%s(%s%d:%s%d)" % (fn, L, first, L, last)
                fmt = number_format(col.get("format"))
                if fmt:
                    c.number_format = fmt
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif ci == 1:
                c.value = spec.get("total_label", "Total")
                c.alignment = Alignment(horizontal="left", vertical="center")
            c.font = Font(name=THEME["font"], size=11, bold=True)
            c.fill = fill(THEME["total_fill"])
            c.border = thin_border(THEME["header_fill"])
    end_row = total_row or last
    if spec.get("freeze", True):
        ws.freeze_panes = ws.cell(row=hrow + 1, column=int(spec.get("freeze_cols", 0)) + 1)
    if spec.get("autofilter", True) and rows:
        ws.auto_filter.ref = "A%d:%s%d" % (hrow, get_column_letter(len(cols)), last)
    autosize(ws, hrow, end_row)
    for i, col in enumerate(cols, 1):
        if col.get("width"):
            ws.column_dimensions[get_column_letter(i)].width = float(col["width"])
    for k, w in (spec.get("column_widths") or {}).items():
        ws.column_dimensions[col_letter(k, headers)].width = float(w)
    if spec.get("title"):  # el título no debe ensanchar la columna A
        pass
    info = SheetInfo(ws, headers, hrow, first, last, total_row)
    apply_conditional(ws, spec.get("conditional") or [], info)
    apply_validations(ws, spec.get("validations") or [], info)
    if spec.get("notes"):
        n = ws.cell(row=end_row + 2, column=1, value=spec["notes"])
        n.font = Font(name=THEME["font"], size=9, italic=True, color="595959")
    ws.page_setup.orientation = "landscape" if len(cols) > 6 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "%d:%d" % (hrow, hrow)
    return info


def data_range(info, col):
    L = col_letter(col, info.headers)
    return "%s%d:%s%d" % (L, info.first, L, max(info.last, info.first))


def apply_conditional(ws, rules, info):
    for r in rules:
        rng = r.get("range") or data_range(info, r.get("column", "A"))
        if re.fullmatch(r"[A-Za-z]{1,3}", rng):
            rng = data_range(info, rng)
        kind = r.get("rule")
        f = fill(r["fill"]) if r.get("fill") else None
        fnt = Font(color=r["font_color"], bold=bool(r.get("bold"))) if (r.get("font_color") or r.get("bold")) else None
        if kind in ("greater_than", "less_than", "equal", "not_equal", "greater_equal", "less_equal"):
            op = {"greater_than": "greaterThan", "less_than": "lessThan", "equal": "equal",
                  "not_equal": "notEqual", "greater_equal": "greaterThanOrEqual", "less_equal": "lessThanOrEqual"}[kind]
            ws.conditional_formatting.add(rng, CellIsRule(operator=op, formula=[str(r["value"])], fill=f or fill("FFC7CE"), font=fnt))
        elif kind == "between":
            ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=[str(r["value"]), str(r["value2"])],
                                                          fill=f or fill("FFEB9C"), font=fnt))
        elif kind == "text_contains":
            first = rng.split(":")[0]
            ws.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("%s",%s))' % (r["value"], first)],
                                                           fill=f or fill("FFEB9C"), font=fnt))
        elif kind == "color_scale":
            ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color=r.get("low", "F8696B"),
                                                              mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                                              end_type="max", end_color=r.get("high", "63BE7B")))
        elif kind == "data_bar":
            ws.conditional_formatting.add(rng, DataBarRule(start_type="min", end_type="max", color=r.get("color", "5B9BD5")))
        else:
            raise SpecError("regla condicional %r no válida (greater_than, less_than, equal, not_equal, between, text_contains, color_scale, data_bar)" % kind)


def apply_validations(ws, vals, info):
    for v in vals:
        L = col_letter(v.get("column", "A"), info.headers)
        end = info.last + int(v.get("extra_rows", 0))
        rng = v.get("range") or "%s%d:%s%d" % (L, info.first, L, max(end, info.first))
        kind = v.get("type", "list")
        if kind == "list":
            items = v.get("values") or []
            formula = '"%s"' % ",".join(str(i).replace('"', "") for i in items)
            if len(formula) > 255:
                raise SpecError("la lista de validación supera 255 caracteres; ponla en una hoja y usa 'source'")
            dv = DataValidation(type="list", formula1=v.get("source") or formula, allow_blank=True)
        elif kind in ("whole", "decimal", "date", "textLength"):
            dv = DataValidation(type=kind, operator="between", formula1=str(v.get("min", 0)),
                                formula2=str(v.get("max", 1000000)), allow_blank=True)
        else:
            raise SpecError("validación %r no válida (list, whole, decimal, date, textLength)" % kind)
        dv.error = v.get("message", "Valor no permitido")
        dv.errorTitle = "Dato inválido"
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(rng)


def resolve_charts(wb, target_ws, charts, infos, own_info=None):
    for ch in charts:
        info = infos.get(ch.get("data_sheet")) if ch.get("data_sheet") else own_info
        if info is None:
            raise SpecError("el gráfico %r necesita 'data_sheet' (hoja con los datos)" % ch.get("title"))
        cat = col_letter(ch.get("categories", "A"), info.headers)
        from openpyxl.utils import column_index_from_string as idx
        series = [col_letter(s, info.headers) for s in ch.get("series", [])]
        if not series:
            raise SpecError("el gráfico %r no tiene 'series'" % ch.get("title"))
        add_chart(target_ws, ch, info.ws, (idx(cat), info.first, info.last),
                  [(idx(s), info.header_row, info.last) for s in series],
                  default_anchor="%s%d" % (get_column_letter(len(info.headers) + 2), info.header_row))


def expand_tokens(formula, infos):
    def rep(m):
        kind, sheet, col = m.group(1), m.group(2), m.group(3)
        if sheet not in infos:
            raise SpecError("la fórmula usa la hoja %r que no existe" % sheet)
        info = infos[sheet]
        L = col_letter(col, info.headers)
        q = "'%s'" % sheet.replace("'", "''")
        if kind == "total":
            if not info.total_row:
                raise SpecError("la hoja %r no tiene fila de total (activa 'total_row')" % sheet)
            return "%s!%s%d" % (q, L, info.total_row)
        return "%s!%s%d:%s%d" % (q, L, info.first, L, info.last)
    return re.sub(r"\{(total|col):([^:}]+):([^}]+)\}", rep, formula)


def build_summary(wb, spec, infos):
    ws = wb.create_sheet(safe_title(spec.get("name", "Resumen")), 0)
    ws.sheet_view.showGridLines = False
    if spec.get("tab_color"):
        ws.sheet_properties.tabColor = spec["tab_color"]
    r = write_title(ws, {"title": spec.get("title", "Resumen"), "subtitle": spec.get("subtitle")}, 2)
    style_header_cell(ws.cell(row=r, column=1, value="Indicador"))
    style_header_cell(ws.cell(row=r, column=2, value="Valor"))
    for k in spec.get("kpis", []):
        r += 1
        a = ws.cell(row=r, column=1, value=k.get("label"))
        a.font = Font(name=THEME["font"], size=11)
        a.border = thin_border()
        v = k.get("formula", k.get("value"))
        if is_formula(v):
            v = add_xlfn_prefix(expand_tokens(v, infos))
        b = ws.cell(row=r, column=2, value=v)
        b.font = Font(name=THEME["font"], size=12, bold=True, color=THEME["title_font"])
        b.border = thin_border()
        b.alignment = Alignment(horizontal="right")
        fmt = number_format(k.get("format"))
        if fmt:
            b.number_format = fmt
    autosize(ws, 1, r, min_w=14)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 28)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 18)
    resolve_charts(wb, ws, spec.get("charts") or [], infos)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def build(spec):
    wb = Workbook()
    wb.remove(wb.active)
    infos = {}
    if not spec.get("sheets"):
        raise SpecError("falta 'sheets' (lista de hojas) en la especificación")
    for s in spec["sheets"]:
        infos[safe_title(s.get("name", "Hoja"))] = build_table(wb, s)
    for s in spec["sheets"]:
        nm = safe_title(s.get("name", "Hoja"))
        resolve_charts(wb, infos[nm].ws, s.get("charts") or [], infos, infos[nm])
    if spec.get("summary"):
        build_summary(wb, spec["summary"], infos)
    p = spec.get("properties") or {}
    wb.properties.title = p.get("title")
    wb.properties.creator = p.get("author", "Lazunox")
    wb.active = 0
    return wb


def main(argv=None):
    ap = argparse.ArgumentParser(description="Crea un .xlsx profesional desde un JSON.", epilog=HELP_SCHEMA,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", nargs="?", help="archivo JSON con la especificación")
    ap.add_argument("-o", "--output", help="archivo .xlsx de salida (sustituye a 'output' del JSON)")
    ap.add_argument("--example", action="store_true", help="imprime un JSON de ejemplo y termina")
    ap.add_argument("--force", action="store_true", help="permite sobrescribir si el archivo ya existe")
    a = ap.parse_args(argv)
    if a.example:
        print(json.dumps(EXAMPLE, ensure_ascii=False, indent=2))
        return 0
    if not a.spec:
        ap.error("falta el archivo JSON de especificación (o usa --example)")
    try:
        spec = load_spec(a.spec)
        out = a.output or spec.get("output")
        if not out:
            raise SpecError("falta el archivo de salida: usa -o archivo.xlsx o 'output' en el JSON")
        if not out.lower().endswith(".xlsx"):
            raise SpecError("la salida debe terminar en .xlsx: %s" % out)
        if os.path.exists(out) and not a.force:
            raise SpecError("ya existe %s; elige un nombre nuevo (p. ej. -v2) o usa --force" % out)
        wb = build(spec)
        wb.save(out)
    except SpecError as e:
        print("Error en la especificación: %s" % e, file=sys.stderr)
        return 2
    except Exception as e:  # noqa: BLE001
        print("Error al crear el Excel: %s: %s" % (type(e).__name__, e), file=sys.stderr)
        return 1
    n = len(wb.sheetnames)
    print("Creado %s (%d hoja%s: %s)" % (out, n, "" if n == 1 else "s", ", ".join(wb.sheetnames)))
    print("Siguiente paso: python recalc_check.py %s  (calcula y verifica las fórmulas)" % out)
    return 0


if __name__ == "__main__":
    sys.exit(main())

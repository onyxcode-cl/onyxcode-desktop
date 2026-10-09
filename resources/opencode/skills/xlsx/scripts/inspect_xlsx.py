#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Vuelca el contenido de un .xlsx: hojas, rango usado, fórmulas, formatos, gráficos, filtros, validaciones.

Uso:
  python inspect_xlsx.py libro.xlsx                 # resumen de todas las hojas
  python inspect_xlsx.py libro.xlsx --sheet Ventas --cells A1:F20   # volcar celdas (valor | fórmula | formato)
  python inspect_xlsx.py libro.xlsx --json
Salida: 0 ok, 2 archivo ausente o ilegible.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import warnings
import zipfile

warnings.filterwarnings("ignore")
from openpyxl import load_workbook  # noqa: E402

CHART_KIND = {"BarChart": "barras/columnas", "LineChart": "líneas", "PieChart": "torta", "AreaChart": "áreas",
              "ScatterChart": "dispersión", "DoughnutChart": "anillo"}


def chart_title(ch):
    try:
        t = ch.title
        if t is None:
            return None
        if isinstance(t, str):
            return t
        return "".join(r.t for p in t.tx.rich.p for r in (p.r or []))
    except Exception:  # noqa: BLE001
        return None


def axis_title(ax):
    try:
        return chart_title(ax) if ax is not None and ax.title is not None else None
    except Exception:  # noqa: BLE001
        return None


def describe_chart(ch):
    kind = type(ch).__name__
    d = {"tipo": CHART_KIND.get(kind, kind), "titulo": chart_title(ch)}
    if kind == "BarChart":
        d["orientacion"] = "horizontal" if ch.type == "bar" else "vertical"
    d["eje_x"] = axis_title(getattr(ch, "x_axis", None))
    d["eje_y"] = axis_title(getattr(ch, "y_axis", None))
    d["series"] = []
    for s in ch.series:
        ref = s.val.numRef.f if s.val is not None and s.val.numRef is not None else None
        cat = None
        if s.cat is not None:
            cat = (s.cat.numRef or s.cat.strRef).f if (s.cat.numRef or s.cat.strRef) else None
        name = s.tx.strRef.f if s.tx is not None and s.tx.strRef is not None else None
        d["series"].append({"nombre": name, "valores": ref, "categorias": cat})
    try:
        a = ch.anchor._from
        d["ancla"] = "%s%d" % (chr(65 + a.col) if a.col < 26 else "col%d" % (a.col + 1), a.row + 1)
    except Exception:  # noqa: BLE001
        pass
    return d


def package_features(path):
    try:
        names = zipfile.ZipFile(path).namelist()
    except zipfile.BadZipFile:
        return ["el archivo no es un zip válido"]
    f = []
    if any("pivotTable" in n for n in names):
        f.append("tablas dinámicas (se conservan solo parcialmente al editar con openpyxl)")
    if any("vbaProject" in n for n in names):
        f.append("macros VBA (xlsm)")
    if any("slicer" in n.lower() for n in names):
        f.append("segmentaciones")
    if any("externalLink" in n for n in names):
        f.append("vínculos a libros externos")
    if any(n.startswith("xl/media/") for n in names):
        f.append("imágenes")
    return f


def inspect(path, only_sheet=None, max_formulas=25):
    wb = load_workbook(path)
    info = {"archivo": path, "hojas": [], "nombres_definidos": [], "caracteristicas": package_features(path)}
    try:
        info["nombres_definidos"] = [n for n in wb.defined_names]
    except Exception:  # noqa: BLE001
        pass
    for ws in wb.worksheets:
        if only_sheet and ws.title != only_sheet:
            continue
        formulas, fmts, nformulas = [], {}, 0
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    nformulas += 1
                    if len(formulas) < max_formulas:
                        formulas.append("%s: %s" % (c.coordinate, v))
                if v is not None and c.number_format != "General":
                    fmts.setdefault(c.column_letter, set()).add(c.number_format)
        widths = {k: round(d.width, 1) for k, d in ws.column_dimensions.items() if d.width}
        h = {
            "nombre": ws.title,
            "estado": ws.sheet_state,
            "rango_usado": ws.dimensions,
            "filas": ws.max_row, "columnas": ws.max_column,
            "paneles_congelados": ws.freeze_panes,
            "autofiltro": ws.auto_filter.ref,
            "celdas_combinadas": [str(r) for r in ws.merged_cells.ranges][:20],
            "anchos_columna": widths,
            "n_formulas": nformulas,
            "formulas": formulas,
            "formatos_numero": {k: sorted(v) for k, v in sorted(fmts.items())},
            "formato_condicional": [str(r.sqref) for r in ws.conditional_formatting],
            "validaciones": [{"rango": str(v.sqref), "tipo": v.type, "formula": v.formula1}
                             for v in ws.data_validations.dataValidation],
            "tablas": {k: t.ref for k, t in ws.tables.items()},
            "graficos": [describe_chart(c) for c in ws._charts],
            "imagenes": len(getattr(ws, "_images", [])),
        }
        info["hojas"].append(h)
    return info


def dump_cells(path, sheet, rng):
    wb = load_workbook(path)
    wv = load_workbook(path, data_only=True)
    ws, wsv = wb[sheet], wv[sheet]
    out = []
    for row in ws[rng]:
        for c in (row if isinstance(row, tuple) else (row,)):
            if c.value is None:
                continue
            cached = wsv[c.coordinate].value
            out.append({"celda": c.coordinate, "valor": c.value, "calculado": cached if str(c.value).startswith("=") else None,
                        "formato": c.number_format, "negrita": bool(c.font and c.font.b)})
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description="Inspecciona un .xlsx (hojas, fórmulas, formatos, gráficos).")
    ap.add_argument("archivo")
    ap.add_argument("--sheet", help="limitar a una hoja")
    ap.add_argument("--cells", help="rango a volcar celda por celda, p. ej. A1:D10 (requiere --sheet)")
    ap.add_argument("--max-formulas", type=int, default=25)
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    if not os.path.isfile(a.archivo):
        print("Error: el archivo no existe: %s" % a.archivo, file=sys.stderr)
        return 2
    try:
        if a.cells:
            if not a.sheet:
                print("Error: --cells requiere --sheet NOMBRE", file=sys.stderr)
                return 2
            cells = dump_cells(a.archivo, a.sheet, a.cells)
            if a.json:
                print(json.dumps(cells, ensure_ascii=False, indent=2, default=str))
            else:
                for c in cells:
                    extra = "" if c["calculado"] is None else "  [guardado: %s]" % c["calculado"]
                    print("%s = %r  (formato %s)%s" % (c["celda"], c["valor"], c["formato"], extra))
            return 0
        info = inspect(a.archivo, a.sheet, a.max_formulas)
    except KeyError:
        print("Error: no existe la hoja %r" % a.sheet, file=sys.stderr)
        return 2
    except Exception as e:  # noqa: BLE001
        print("Error: no se pudo abrir el libro (%s: %s)" % (type(e).__name__, e), file=sys.stderr)
        return 2
    if a.json:
        print(json.dumps(info, ensure_ascii=False, indent=2, default=str))
        return 0
    print("Archivo: %s | hojas: %d" % (info["archivo"], len(info["hojas"])))
    if info["caracteristicas"]:
        print("Contiene (cuidado al editar): " + "; ".join(info["caracteristicas"]))
    for h in info["hojas"]:
        print("\n== Hoja %r (%s) rango %s [%d filas x %d columnas]" % (h["nombre"], h["estado"], h["rango_usado"], h["filas"], h["columnas"]))
        print("  paneles congelados: %s | autofiltro: %s | combinadas: %s" % (h["paneles_congelados"], h["autofiltro"], h["celdas_combinadas"] or "no"))
        print("  anchos: %s" % (h["anchos_columna"] or "por defecto"))
        print("  formatos de número: %s" % (h["formatos_numero"] or "ninguno"))
        print("  fórmulas: %d" % h["n_formulas"])
        for f in h["formulas"]:
            print("    " + f)
        if h["formato_condicional"]:
            print("  formato condicional: %s" % h["formato_condicional"])
        if h["validaciones"]:
            print("  validaciones: %s" % h["validaciones"])
        for g in h["graficos"]:
            print("  gráfico: %s" % json.dumps(g, ensure_ascii=False))
        if h["imagenes"]:
            print("  imágenes: %d" % h["imagenes"])
    return 0


if __name__ == "__main__":
    sys.exit(main())

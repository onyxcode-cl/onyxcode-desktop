#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verificación general de un .xlsx: abre, hojas, errores de fórmula, columnas estrechas, números como texto.

Uso:  python check_xlsx.py libro.xlsx [--render DIR] [--json]
  --render DIR  además renderiza a PNG con _shared/render.py (necesita GridCraft, de ArtCraft) para mirar las imágenes.
Salida: 0 sin problemas, 1 hay problemas, 2 no se pudo abrir.
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import os
import re
import subprocess
import sys
import warnings
import zipfile

warnings.filterwarnings("ignore")
from openpyxl import load_workbook  # noqa: E402

ERRS = ("#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!")
NUM_TEXT = re.compile(r"^\s*[-+]?\$?\s*\d{1,3}([.,]\d{3})*([.,]\d+)?\s*%?\s*$|^\s*[-+]?\$?\d+([.,]\d+)?\s*%?\s*$")


def est_len(v, fmt):
    if isinstance(v, (_dt.datetime, _dt.date)):
        return 10
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        dec = 2 if "0.00" in str(fmt) else (1 if "0.0" in str(fmt) else 0)
        return len("{:,.{d}f}".format(v, d=dec)) + (2 if "$" in str(fmt) or "%" in str(fmt) else 0)
    return max((len(x) for x in str(v).split("\n")), default=0)


def check(path):
    out = {"archivo": path, "hojas": [], "problemas": [], "avisos": []}
    try:
        zipfile.ZipFile(path).testzip()
        wb = load_workbook(path)
        wv = load_workbook(path, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise RuntimeError("no se pudo abrir el libro: %s: %s" % (type(e).__name__, e))
    if not wb.sheetnames:
        out["problemas"].append("el libro no tiene hojas")
    uncached = 0
    for ws in wb.worksheets:
        out["hojas"].append({"nombre": ws.title, "rango": ws.dimensions, "graficos": len(ws._charts)})
        if ws.max_row <= 1 and ws.max_column <= 1 and ws["A1"].value is None:
            out["avisos"].append("la hoja %r está vacía" % ws.title)
        wsv = wv[ws.title]
        narrow, text_nums, nforms = [], [], 0
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith("=") and len(v) > 1:
                    nforms += 1
                    cv = wsv[c.coordinate].value
                    if isinstance(cv, str) and cv in ERRS:
                        out["problemas"].append("%s!%s: la fórmula %s da %s" % (ws.title, c.coordinate, v, cv))
                    if cv is None:
                        uncached += 1
                    if "#REF!" in v:
                        out["problemas"].append("%s!%s: la fórmula contiene #REF!" % (ws.title, c.coordinate))
                    continue
                if isinstance(v, str):
                    if v in ERRS:
                        out["problemas"].append("%s!%s: valor de error %s" % (ws.title, c.coordinate, v))
                    elif NUM_TEXT.match(v) and c.number_format in ("General", "@") and len(text_nums) < 200:
                        text_nums.append(c.coordinate)
                w = ws.column_dimensions[c.column_letter].width or 8.43
                if not (isinstance(v, str) and c.alignment.wrap_text) and ws.column_dimensions[c.column_letter].width:
                    if (isinstance(v, (int, float, _dt.date)) or True) and est_len(v, c.number_format) > w * 1.15 \
                            and not isinstance(v, str):
                        narrow.append(c.coordinate)
        if narrow:
            out["problemas"].append("hoja %r: %d celda(s) numéricas/fecha no caben en su columna (se verían ####): %s"
                                    % (ws.title, len(narrow), ", ".join(narrow[:8])))
        if text_nums:
            out["problemas"].append("hoja %r: %d número(s) guardado(s) como texto (no suman ni ordenan bien): %s"
                                    % (ws.title, len(text_nums), ", ".join(text_nums[:8])))
        # texto largo en columna estrecha sin ajuste y con vecina ocupada
        for col in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)):
            L = col[0].column_letter
            w = ws.column_dimensions[L].width
            if not w:
                continue
            cut = [c.coordinate for c in col if isinstance(c.value, str) and not c.value.startswith("=")
                   and not c.alignment.wrap_text and len(c.value) > w * 1.3 and c.row > 1
                   and ws.cell(row=c.row, column=c.column + 1).value is not None]
            if cut:
                out["avisos"].append("hoja %r: texto cortado por la columna vecina en %s (ensancha la columna %s o activa ajuste de texto)"
                                     % (ws.title, ", ".join(cut[:5]), L))
    if uncached:
        out["avisos"].append("%d fórmula(s) sin valor guardado (normal en archivos de openpyxl): ejecuta recalc_check.py para "
                             "calcularlas y verificar errores" % uncached)
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description="Verifica un .xlsx (estructura, errores, anchos, números como texto).")
    ap.add_argument("archivo")
    ap.add_argument("--render", metavar="DIR", help="renderiza a PDF/PNG en DIR usando _shared/render.py")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    if not os.path.isfile(a.archivo):
        print("Error: el archivo no existe: %s" % a.archivo, file=sys.stderr)
        return 2
    try:
        res = check(a.archivo)
    except RuntimeError as e:
        print("Error: %s" % e, file=sys.stderr)
        return 2
    if a.render:
        here = os.path.dirname(os.path.abspath(__file__))
        render = os.path.join(here, "..", "..", "_shared", "render.py")
        if not os.path.isfile(render):
            res["avisos"].append("no encuentro _shared/render.py: no se pudo renderizar")
        else:
            os.makedirs(a.render, exist_ok=True)
            r = subprocess.run([sys.executable, render, a.archivo, "--outdir", a.render], capture_output=True, text=True)
            try:
                res["render"] = json.loads(r.stdout)
                res["avisos"].extend(res["render"].get("avisos", []))
                if not res["render"].get("ok"):
                    res["avisos"].append("el render falló o ArtCraft (GridCraft) no está disponible: la verificación visual no se hizo")
            except ValueError:
                res["avisos"].append("render.py no devolvió JSON: %s" % (r.stderr or r.stdout)[:200])
    res["ok"] = not res["problemas"]
    if a.json:
        print(json.dumps(res, ensure_ascii=False, indent=2))
        return 0 if res["ok"] else 1
    print("Archivo: %s abre correctamente. Hojas: %s" % (res["archivo"], ", ".join("%s (%s, %d gráf.)" % (h["nombre"], h["rango"], h["graficos"]) for h in res["hojas"])))
    for p in res["problemas"]:
        print("PROBLEMA: " + p)
    for v in res["avisos"]:
        print("AVISO: " + v)
    if res.get("render", {}).get("pngs"):
        print("Imágenes para revisar:")
        for p in res["render"]["pngs"]:
            print("  " + p)
    print("RESULTADO: " + ("sin problemas" if res["ok"] else "hay %d problema(s)" % len(res["problemas"])))
    return 0 if res["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())

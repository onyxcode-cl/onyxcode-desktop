#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Calcula y verifica las fórmulas de un .xlsx (openpyxl NO calcula fórmulas).

Recalcula con GridCraft (programa de ArtCraft, que instala la app) y lee los valores. Detecta #DIV/0!, #REF!,
#NAME?, #VALUE!, #N/A, #NUM!, #NULL! y fórmulas que apuntan a celdas vacías. Sin GridCraft (o si no da un resultado
fiable) lo dice claramente y hace solo una verificación sintáctica, con `motor` null.

Uso:  python recalc_check.py libro.xlsx [--json] [--show N] [--workdir DIR]
Salida: 0 sin problemas, 1 hay errores/advertencias graves, 2 uso incorrecto.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import tempfile

import warnings

warnings.filterwarnings("ignore")
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import column_index_from_string, range_boundaries
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "_shared"))
import render  # noqa: E402  (AVISO_SIN_ARTCRAFT)
import render_grid  # noqa: E402  GridCraft (ArtCraft)

ERRORS = ("#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!")
KNOWN_FUNCS = set("""ABS AND AVERAGE AVERAGEIF AVERAGEIFS CEILING CHOOSE COLUMN COLUMNS CONCAT CONCATENATE COUNT COUNTA
COUNTBLANK COUNTIF COUNTIFS DATE DATEDIF DAY DAYS EDATE EOMONTH EXP FILTER FIND FLOOR HLOOKUP HOUR IF IFERROR IFNA IFS
INDEX INDIRECT INT ISBLANK ISERROR ISNA ISNUMBER ISTEXT LARGE LEFT LEN LN LOG LOG10 LOOKUP LOWER MATCH MAX MAXIFS MEDIAN
MID MIN MINIFS MINUTE MOD MONTH NETWORKDAYS NOT NOW OFFSET OR PI PMT POWER PRODUCT PROPER RANK RIGHT ROUND ROUNDDOWN
ROUNDUP ROW ROWS SEARCH SECOND SMALL SORT SQRT STDEV SUBSTITUTE SUBTOTAL SUM SUMIF SUMIFS SUMPRODUCT SWITCH TEXT TEXTJOIN
TODAY TRIM TRUE FALSE UNIQUE UPPER VALUE VLOOKUP WEEKDAY XLOOKUP XOR YEAR NPV IRR FV PV RATE VAR TRANSPOSE COUNTUNIQUE
ISERR ISEVEN ISODD N NA T TRUNC SIGN ROUNDUP REPT CHAR CODE CLEAN EXACT FIXED DOLLAR""".split())


def cells_in_ref(ref):
    ref = ref.replace("$", "")
    try:
        c1, r1, c2, r2 = range_boundaries(ref)
    except Exception:  # noqa: BLE001
        return None
    if None in (c1, r1, c2, r2) or (c2 - c1 + 1) * (r2 - r1 + 1) > 5000:
        return None
    return c1, r1, c2, r2


REF_RE = re.compile(r"(?:(?:'([^']+)'|([A-Za-z_][\w\.]*))!)?(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)(?![\w(])")


def empty_refs(formula, wb, own_sheet):
    """Referencias (celda única o rango completo) cuyo contenido está vacío."""
    body = re.sub(r'"[^"]*"', '""', formula)
    found = []
    for m in REF_RE.finditer(body):
        sheet = m.group(1) or m.group(2) or own_sheet
        if sheet not in wb.sheetnames:
            continue
        box = cells_in_ref(m.group(3))
        if not box:
            continue
        c1, r1, c2, r2 = box
        ws = wb[sheet]
        if all(ws.cell(row=r, column=c).value is None for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)):
            found.append("%s!%s" % (sheet, m.group(3).replace("$", "")))
    return found


def syntax_issues(formula, wb):
    out = []
    body = re.sub(r'"[^"]*"', '""', formula)
    if formula.count('"') % 2:
        out.append("comillas sin cerrar")
    if body.count("(") != body.count(")"):
        out.append("paréntesis desbalanceados")
    for fn in re.findall(r"([A-Za-z_][A-Za-z0-9_\.]*)\(", body):
        u = fn.upper().replace("_XLFN.", "")
        if u not in KNOWN_FUNCS:
            out.append("función no reconocida %s (¿nombre en español o mal escrito? en el archivo van en inglés)" % fn)
    for m in re.finditer(r"(?:'([^']+)'|([A-Za-z_][\w\.]*))!", body):
        s = m.group(1) or m.group(2)
        if s not in wb.sheetnames:
            out.append("referencia a la hoja inexistente %r" % s)
    if "#REF!" in formula:
        out.append("contiene #REF!")
    return out


def analyse(path, show):
    wbf = load_workbook(path)  # fórmulas
    formulas = []
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and len(c.value) > 1:
                    formulas.append((ws.title, c.coordinate, c.value))
    res = {"archivo": path, "motor": None, "formulas": len(formulas), "errores": [],
           "vacias": [], "sintaxis": [], "valores": [], "avisos": []}
    grid = render_grid.recalc(path, res["avisos"])
    if grid is not None:
        res["motor"] = "gridcraft"
        mapa = {(v["hoja"], v["celda"]): v["valor"] for v in grid["valores"]}
    else:
        mapa = None
        if render_grid.find_gridcraft() is None:
            res["avisos"].append(render.AVISO_SIN_ARTCRAFT)
        res["avisos"].append("Solo verificación sintáctica: NO se calcularon los valores. Dile al usuario que los valores "
                             "no están verificados.")
    for sheet, coord, f in formulas:
        for p in syntax_issues(f, wbf):
            res["sintaxis"].append({"hoja": sheet, "celda": coord, "formula": f, "problema": p})
        for e in empty_refs(f, wbf, sheet):
            res["vacias"].append({"hoja": sheet, "celda": coord, "formula": f, "referencia": e})
        if mapa is not None:
            v = mapa.get((sheet, coord))
            if isinstance(v, str) and (v in ERRORS or v.startswith("Err:")):
                res["errores"].append({"hoja": sheet, "celda": coord, "formula": f, "valor": v})
            if len(res["valores"]) < show:
                res["valores"].append({"hoja": sheet, "celda": coord, "formula": f, "valor": v})
    return res


def main(argv=None):
    ap = argparse.ArgumentParser(description="Recalcula (GridCraft) y verifica las fórmulas de un .xlsx.")
    ap.add_argument("archivo", help="libro .xlsx a verificar (no se modifica)")
    ap.add_argument("--json", action="store_true", help="salida en JSON")
    ap.add_argument("--show", type=int, default=15, help="cuántos valores calculados mostrar (def. 15)")
    ap.add_argument("--workdir", help="carpeta de trabajo temporal (def.: una carpeta temporal del sistema)")
    a = ap.parse_args(argv)
    if not os.path.isfile(a.archivo):
        print("Error: el archivo no existe: %s" % a.archivo, file=sys.stderr)
        return 2
    base = a.workdir or tempfile.mkdtemp(prefix="xlsx-recalc-")
    os.makedirs(base, exist_ok=True)
    try:
        res = analyse(a.archivo, a.show)
    except Exception as e:  # noqa: BLE001
        print("Error al abrir el libro: %s: %s" % (type(e).__name__, e), file=sys.stderr)
        return 2
    finally:
        if not a.workdir:
            shutil.rmtree(base, ignore_errors=True)
    bad = bool(res["errores"] or res["sintaxis"] or res["vacias"])
    res["ok"] = not bad
    if a.json:
        print(json.dumps(res, ensure_ascii=False, indent=2, default=str))
        return 1 if bad else 0
    print("Archivo: %s" % res["archivo"])
    quien = {"gridcraft": "GridCraft (ArtCraft)"}.get(res["motor"], "NADA (solo verificación sintáctica)")
    print("Fórmulas: %d | Recalculado con: %s" % (res["formulas"], quien))
    for av in res["avisos"]:
        print("AVISO: " + av)
    for e in res["errores"]:
        print("ERROR  %s!%s  %s  ->  %s" % (e["hoja"], e["celda"], e["formula"], e["valor"]))
    for e in res["sintaxis"]:
        print("SINTAXIS  %s!%s  %s  ->  %s" % (e["hoja"], e["celda"], e["formula"], e["problema"]))
    for e in res["vacias"]:
        print("VACÍA  %s!%s  %s  apunta a celdas vacías: %s" % (e["hoja"], e["celda"], e["formula"], e["referencia"]))
    if res["valores"]:
        print("Valores calculados (primeras %d):" % len(res["valores"]))
        for v in res["valores"]:
            print("  %s!%s  %s = %s" % (v["hoja"], v["celda"], v["formula"], v["valor"]))
    print("RESULTADO: " + ("sin problemas detectados" if not bad else "HAY PROBLEMAS: corrígelos y vuelve a ejecutar"))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())

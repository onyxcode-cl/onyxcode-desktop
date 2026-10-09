# -*- coding: utf-8 -*-
"""Utilidades compartidas por los scripts de la skill xlsx (solo openpyxl + stdlib)."""
from __future__ import annotations

import datetime as _dt
import re

from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

# Formatos de número. Excel guarda el código en inglés y lo muestra según la
# configuración regional del usuario (en es-CL: 1.234,5 y 12,5 %).
FORMATS = {
    "text": "@",
    "general": "General",
    "integer": "#,##0",
    "number": "#,##0.00",
    "decimal": "#,##0.00",
    "clp": '"$"#,##0',
    "currency": '"$"#,##0',
    "usd": '"US$"#,##0.00',
    "eur": '"€"#,##0.00',
    "uf": '"UF "#,##0.00',
    "percent": "0.0%",
    "percent0": "0%",
    "date": "dd-mm-yyyy",
    "date_long": 'd "de" mmmm "de" yyyy',
    "datetime": "dd-mm-yyyy hh:mm",
    "month": "mmm-yyyy",
}

NUMERIC_KINDS = {"integer", "number", "decimal", "clp", "currency", "usd", "eur", "uf", "percent", "percent0"}
DATE_KINDS = {"date", "date_long", "datetime", "month"}

THEME = {
    "header_fill": "1F3864",
    "header_font": "FFFFFF",
    "title_font": "1F3864",
    "zebra": "F2F6FC",
    "border": "BFC8D6",
    "total_fill": "DCE6F4",
    "font": "Calibri",
}

CHART_TYPES = {"bar": ("bar", None), "column": ("col", None), "col": ("col", None),
               "line": ("line", None), "pie": ("pie", None),
               "stacked_column": ("col", "stacked"), "stacked_bar": ("bar", "stacked")}


class SpecError(ValueError):
    """Error accionable en la especificación del usuario."""


def number_format(kind):
    if kind is None:
        return None
    return FORMATS.get(str(kind).lower(), kind)  # permite código personalizado


def is_formula(v):
    return isinstance(v, str) and v.startswith("=") and len(v) > 1


def parse_date(v):
    """Convierte 'YYYY-MM-DD' / 'DD-MM-YYYY' / 'DD/MM/YYYY' (y con hora) a datetime; si no, devuelve v."""
    if isinstance(v, (_dt.date, _dt.datetime)) or not isinstance(v, str):
        return v
    s = v.strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return _dt.datetime.strptime(s, fmt)
        except ValueError:
            pass
    return v


def col_letter(ref, headers=None, ws=None, header_row=1):
    """Acepta 'B', 2 o el texto de un encabezado y devuelve la letra de columna."""
    if isinstance(ref, int):
        return get_column_letter(ref)
    s = str(ref).strip()
    if re.fullmatch(r"[A-Za-z]{1,3}", s):
        # ambiguo con un encabezado corto: priorizar encabezado exacto si hay lista
        if headers and s in headers:
            return get_column_letter(headers.index(s) + 1)
        return s.upper()
    if headers and s in headers:
        return get_column_letter(headers.index(s) + 1)
    if ws is not None:
        for c in ws[header_row]:
            if c.value is not None and str(c.value).strip() == s:
                return c.column_letter
    raise SpecError("no encuentro la columna %r (usa la letra, p. ej. 'B', o el texto exacto del encabezado)" % s)


def thin_border(color=None):
    side = Side(style="thin", color=color or THEME["border"])
    return Border(left=side, right=side, top=side, bottom=side)


def style_header_cell(c):
    c.font = Font(name=THEME["font"], bold=True, color=THEME["header_font"], size=11)
    c.fill = PatternFill("solid", start_color=THEME["header_fill"], end_color=THEME["header_fill"])
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border(THEME["header_fill"])


def display_len(v, fmt=None):
    """Largo aproximado del texto visible de un valor (para anchos automáticos)."""
    if v is None:
        return 0
    if is_formula(v):
        return 12
    if isinstance(v, bool):
        return 5
    if isinstance(v, (_dt.datetime, _dt.date)):
        return 16 if fmt and "h" in str(fmt) else (18 if fmt and "mmmm" in str(fmt) else 11)
    if isinstance(v, (int, float)):
        f = str(fmt or "")
        dec = 2 if ("0.00" in f) else (1 if "0.0" in f else 0)
        txt = "{:,.{d}f}".format(v, d=dec)
        extra = 4 if ("US$" in f or "UF" in f or "€" in f) else (1 if ("$" in f or "%" in f) else 0)
        return len(txt) + extra + 1
    return max(len(line) for line in str(v).split("\n"))


def autosize(ws, first_row, last_row, min_w=8, max_w=60, kinds=None):
    for cells in ws.iter_cols(min_row=first_row, max_row=last_row):
        letter = cells[0].column_letter
        best = 0
        for c in cells:
            n = display_len(c.value, c.number_format)
            if c.row == first_row:
                n = n * 1.1 + 2  # encabezados en negrita + margen para el filtro
            best = max(best, n)
        ws.column_dimensions[letter].width = max(min_w, min(max_w, round(best + 2, 1)))


def add_chart(target_ws, spec, data_ws, cat_ref, series_refs, default_anchor="H2"):
    """Crea un gráfico y lo añade a target_ws.

    cat_ref: (col_idx, first_row, last_row) de las categorías (sin encabezado).
    series_refs: lista de (col_idx, header_row, last_row); la celda de encabezado da el nombre.
    spec: type, title, x_title, y_title, anchor, width, height (cm), style, data_labels, legend.
    """
    kind = str(spec.get("type", "column")).lower()
    if kind not in CHART_TYPES:
        raise SpecError("tipo de gráfico %r no válido; usa: bar, column, line, pie, stacked_column, stacked_bar" % kind)
    base, grouping = CHART_TYPES[kind]
    if base == "line":
        ch = LineChart()
    elif base == "pie":
        ch = PieChart()
    else:
        ch = BarChart()
        ch.type = base  # "col" vertical, "bar" horizontal
        if grouping:
            ch.grouping = "stacked"
            ch.overlap = 100
    ch.title = spec.get("title")
    ch.width = float(spec.get("width", 16))
    ch.height = float(spec.get("height", 8.5))
    if spec.get("style"):
        ch.style = int(spec["style"])
    cc, cr1, cr2 = cat_ref
    cats = Reference(data_ws, min_col=cc, min_row=cr1, max_row=cr2)
    if base == "pie":
        series_refs = series_refs[:1]
    for (c, hr, lr) in series_refs:
        ref = Reference(data_ws, min_col=c, min_row=hr, max_row=lr)
        ch.add_data(ref, titles_from_data=True)
    ch.set_categories(cats)
    if base == "line":
        for s in ch.series:
            s.smooth = False
            s.marker.symbol = "circle"
    if base != "pie":
        ch.x_axis.title = spec.get("x_title")
        ch.y_axis.title = spec.get("y_title")
        # openpyxl >= 3.1 oculta los ejes salvo que se pida explícitamente
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        if kind in ("column", "col", "stacked_column") and len(series_refs) == 1 and spec.get("legend") is None:
            ch.legend = None
    if spec.get("legend") is False:
        ch.legend = None
    elif spec.get("legend") in ("right", "bottom", "top", "left") and ch.legend is not None:
        ch.legend.position = {"right": "r", "bottom": "b", "top": "t", "left": "l"}[spec["legend"]]
    if spec.get("data_labels") or base == "pie":
        ch.dataLabels = DataLabelList()
        if base == "pie":
            ch.dataLabels.showPercent = True
            ch.dataLabels.showVal = False
        else:
            ch.dataLabels.showVal = True
        ch.dataLabels.showSerName = False
        ch.dataLabels.showCatName = False
        ch.dataLabels.showLegendKey = False
    target_ws.add_chart(ch, spec.get("anchor", default_anchor))
    return ch


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)

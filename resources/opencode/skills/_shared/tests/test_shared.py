"""Pruebas de humo de _shared (unittest). Ejecutar: python -m unittest tests/test_shared.py"""
import json
import os
import subprocess
import sys
import tempfile
import unittest

AQUI = os.path.dirname(os.path.abspath(__file__))
SHARED = os.path.dirname(AQUI)
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def correr(script, *args):
    r = subprocess.run([sys.executable, os.path.join(SHARED, script)] + list(args),
                       capture_output=True, text=True, timeout=300)
    return r.returncode, r.stdout, r.stderr


def hay_lo():
    import shutil
    return os.path.isfile(SOFFICE) or shutil.which("soffice") is not None


class PruebasShared(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = cls.tmp.name
        from reportlab.pdfgen import canvas
        cls.pdf = os.path.join(cls.dir, "mini.pdf")
        c = canvas.Canvas(cls.pdf)
        c.setFont("Helvetica", 18)
        for i in range(40):
            c.drawString(100, 780 - i * 16, "Hola Lazunox, linea %d de texto de prueba" % i)
        c.showPage()
        c.showPage()  # segunda página en blanco
        c.save()

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_render_pdf(self):
        out = os.path.join(self.dir, "o1")
        code, so, _ = correr("render.py", self.pdf, "--outdir", out, "--dpi", "50")
        self.assertEqual(code, 0, so)
        j = json.loads(so)
        self.assertTrue(j["ok"])
        self.assertEqual(j["paginas"], 2)
        self.assertEqual(len(j["pngs"]), 2)
        for p in j["pngs"]:
            self.assertTrue(os.path.isfile(p))
        # verify_images: la 2.ª página está en blanco
        code, so, _ = correr("verify_images.py", *j["pngs"])
        self.assertEqual(code, 0, so)
        v = json.loads(so)
        self.assertTrue(any("blanco" in a for a in v["paginas"][1]["avisos"]))
        self.assertFalse(any("blanco" in a for a in v["paginas"][0]["avisos"]))

    def test_max_pages(self):
        out = os.path.join(self.dir, "o2")
        code, so, _ = correr("render.py", self.pdf, "--outdir", out, "--max-pages", "1")
        j = json.loads(so)
        self.assertEqual(len(j["pngs"]), 1)
        self.assertEqual(j["paginas"], 2)
        self.assertTrue(j["avisos"])

    def test_archivo_inexistente_y_uso(self):
        code, so, _ = correr("render.py", os.path.join(self.dir, "no.pdf"), "--outdir", self.dir)
        self.assertEqual(code, 1)
        self.assertFalse(json.loads(so)["ok"])
        code, _, _ = correr("render.py")
        self.assertEqual(code, 2)

    def test_verify_borde(self):
        from PIL import Image, ImageDraw
        p = os.path.join(self.dir, "borde.png")
        im = Image.new("RGB", (200, 300), "white")
        ImageDraw.Draw(im).rectangle([0, 50, 199, 250], fill="black")
        im.save(p)
        code, so, _ = correr("verify_images.py", p)
        v = json.loads(so)
        self.assertTrue(v["paginas"][0]["avisos"])

    def test_doctor(self):
        code, so, _ = correr("doctor.py")
        j = json.loads(so)
        self.assertIn("librerias", j)
        self.assertIn("faltan", j)
        self.assertEqual(code, 0 if not j["faltan"] else 1)

    def test_render_docx(self):
        if not hay_lo():
            self.skipTest("LibreOffice no está instalado")
        import docx
        d = docx.Document()
        d.add_heading("Prueba", 1)
        d.add_paragraph("Contenido de prueba.")
        ruta = os.path.join(self.dir, "mini.docx")
        d.save(ruta)
        out = os.path.join(self.dir, "o3")
        code, so, _ = correr("render.py", ruta, "--outdir", out)
        self.assertEqual(code, 0, so)
        j = json.loads(so)
        self.assertTrue(j["ok"])
        self.assertGreaterEqual(len(j["pngs"]), 1)

    def test_render_docx_ruta_con_espacios(self):
        # Regresión: el perfil de LibreOffice iba en una URL sin escapar y abortaba (código -6) si la ruta tenía espacios.
        if not hay_lo():
            self.skipTest("LibreOffice no está instalado")
        import docx
        d = docx.Document()
        d.add_paragraph("Ruta con espacios.")
        ruta = os.path.join(self.dir, "con espacio.docx")
        d.save(ruta)
        out = os.path.join(self.dir, "carpeta con espacios", "salida")
        code, so, _ = correr("render.py", ruta, "--outdir", out)
        self.assertEqual(code, 0, so)
        self.assertTrue(json.loads(so)["ok"], so)


    def test_buscar_soffice_en_la_carpeta_de_binarios_de_la_app(self):
        # LibreOffice instalado bajo demanda por la app: LAZUNOX_BIN_DIR/libreoffice-<versión>/LibreOffice.app/Contents/MacOS/soffice
        import importlib.util
        spec = importlib.util.spec_from_file_location("render_mod", os.path.join(SHARED, "render.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        with tempfile.TemporaryDirectory() as base:
            exe = os.path.join(base, "libreoffice-9.9.9", "LibreOffice.app", "Contents", "MacOS", "soffice")
            os.makedirs(os.path.dirname(exe))
            with open(exe, "w") as f:
                f.write("#!/bin/sh\n")
            os.chmod(exe, 0o755)
            viejo = {k: os.environ.pop(k, None) for k in ("LAZUNOX_SOFFICE", "LAZUNOX_BIN_DIR")}
            try:
                os.environ["LAZUNOX_BIN_DIR"] = base
                self.assertEqual(mod.buscar_soffice(), exe)
                os.environ["LAZUNOX_SOFFICE"] = exe  # tiene prioridad
                self.assertEqual(mod.buscar_soffice(), exe)
            finally:
                for k, v in viejo.items():
                    os.environ.pop(k, None)
                    if v is not None:
                        os.environ[k] = v

if __name__ == "__main__":
    unittest.main()

"""render_grid (GridCraft para xlsx) con un gridcraft-cli FALSO: búsqueda, prioridad, valores, guardas y avisos sin resultado fiable.
La prueba con el binario real se salta si no se pasa GRIDCRAFT_CLI=/ruta/a/gridcraft-cli.

Ejecutar:  $LAZUNOX_PYTHON -W ignore tests/test_craft_grid.py   (o python -m unittest discover -s tests)
"""
from __future__ import annotations

import hashlib
import json
import os
import shutil
import stat
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

AQUI = os.path.dirname(os.path.abspath(__file__))
SHARED = os.path.dirname(AQUI)
RENDER = os.path.join(SHARED, "render.py")
RECALC = os.path.join(SHARED, "..", "xlsx", "scripts", "recalc_check.py")
sys.path.insert(0, SHARED)
import render_grid  # noqa: E402

REAL = os.environ.get("GRIDCRAFT_CLI")

# Plantilla del gridcraft-cli falso: convert copia valores.json; functions imprime catalogo.json; run copia pdf
# a la ruta de file.exportPdf. Los datos están junto al propio script (el entorno de la app no pasa variables).
FALSO = """#!/bin/sh
D=$(dirname "$0")
[ -f "$D/fallo" ] && exit 1
case "$1" in
  convert) cp "$D/valores.json" "$3" ;;
  functions) cat "$D/catalogo.json" ;;
  run)
    p=$(printf '%s' "$5" | sed 's/.*"path": *"\\([^"]*\\)".*/\\1/')
    cp "$D/libro.pdf" "$p" && echo '{"command":"file.exportPdf","result":{"pages":2}}' ;;
  *) exit 2 ;;
esac
"""


def celda(t, v, formula=None):
    return {"value": {"t": t, "v": v}, "formula": {"text": formula, "array": None} if formula else None, "style": 0}


def libro_json(celdas, hoja="Datos"):
    """celdas: {(fila0, col0): celda}. Mismo esquema que `gridcraft-cli convert libro.xlsx salida.json`."""
    rows = {}
    for (f, c), cel in celdas.items():
        rows.setdefault(str(f), {})[str(c)] = cel
    return {"sheets": [{"name": hoja, "cells": {"len": len(rows), "bands": {"0": {"rows": rows}}}}]}


VALORES = libro_json({
    (0, 0): celda("Number", 5.0),
    (0, 1): celda("Number", 0.0),
    (0, 2): celda("Error", "Div0", "A1/B1"),
    (1, 2): celda("Number", 0.3333333333333333, "1/3"),
    (2, 2): celda("Number", 5.0, "SUM(A1:B1)"),
    (3, 2): celda("Text", "x", '"x"'),
    (4, 2): celda("Bool", True, "ISBLANK(C9)"),
    (5, 2): celda("Error", "NA", "NA()"),
    (6, 2): celda("Error", "Name", "FOO(A1)"),
})
CATALOGO = [{"name": n, "category": "x", "description": "", "signature": ""} for n in ("SUM", "IF", "FILTER", "TEXTJOIN", "ROUND")]


def pdf_de_prueba(ruta):
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(ruta)
    c.drawString(100, 700, "Hoja 1")
    c.showPage()
    c.drawString(100, 700, "Hoja 2")
    c.showPage()
    c.save()


def entorno(**extra):
    """Entorno de prueba sin variables de GridCraft ni de la app; lo que se pase en `extra` sí."""
    env = {k: v for k, v in os.environ.items() if not k.startswith("LAZUNOX_")}
    env.update(extra)
    return env


class FakeCraft(unittest.TestCase):
    """Base: crea una carpeta con un gridcraft-cli falso y sus datos."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="gridcraft-test-")
        cls.pdf = os.path.join(cls.tmp, "libro.pdf")
        pdf_de_prueba(cls.pdf)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def setUp(self):
        self.d = tempfile.mkdtemp(prefix="gridcraft-fake-")
        self.bin = os.path.join(self.d, "bin")
        os.makedirs(self.bin)
        self.fake = self.instalar(self.bin)
        self.xlsx = os.path.join(self.d, "libro.xlsx")
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "Datos"
        wb.active["A1"] = 5
        wb.active["C1"] = "=A1/B1"
        wb.save(self.xlsx)

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    def instalar(self, carpeta, valores=VALORES, catalogo=CATALOGO, fallo=False):
        os.makedirs(carpeta, exist_ok=True)
        ruta = os.path.join(carpeta, "gridcraft-cli")
        with open(ruta, "w") as f:
            f.write(FALSO)
        os.chmod(ruta, os.stat(ruta).st_mode | stat.S_IEXEC)
        with open(os.path.join(carpeta, "valores.json"), "w") as f:
            json.dump(valores, f)
        with open(os.path.join(carpeta, "catalogo.json"), "w") as f:
            json.dump(catalogo, f)
        shutil.copyfile(self.pdf, os.path.join(carpeta, "libro.pdf"))
        if fallo:
            open(os.path.join(carpeta, "fallo"), "w").close()
        return ruta


class Busqueda(FakeCraft):
    def test_variable_del_motor_gana_y_sin_marca_no_se_usa(self):
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=self.fake, LAZUNOX_BIN_DIR=self.bin), clear=True):
            self.assertEqual(render_grid.find_gridcraft(), self.fake)
        otro = self.instalar(os.path.join(self.d, "otro"))
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=otro, LAZUNOX_BIN_DIR=self.bin), clear=True):
            self.assertEqual(render_grid.find_gridcraft(), otro)
        # en la carpeta de la app solo cuenta con la marca de instalación completa
        craft = os.path.join(self.bin, "craft-0.5.0", "grid")
        os.makedirs(craft)
        shutil.copy2(self.fake, os.path.join(craft, "gridcraft-cli"))
        with mock.patch.dict(os.environ, entorno(LAZUNOX_BIN_DIR=self.bin), clear=True):
            self.assertIsNone(render_grid.find_gridcraft())
        open(os.path.join(self.bin, "craft-0.5.0", "installed-gridcraft.json"), "w").write("{}")
        with mock.patch.dict(os.environ, entorno(LAZUNOX_BIN_DIR=self.bin), clear=True):
            self.assertEqual(render_grid.find_gridcraft(), os.path.join(craft, "gridcraft-cli"))

    def test_no_busca_en_el_path_ni_en_rutas_fijas(self):
        con_path = os.path.join(self.d, "solo-path")
        self.instalar(con_path)
        with mock.patch.dict(os.environ, entorno(PATH=con_path), clear=True):
            self.assertIsNone(render_grid.find_gridcraft())

    def test_ruta_que_no_es_ejecutable_se_ignora(self):
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=os.path.join(self.d, "no-existe")), clear=True):
            self.assertIsNone(render_grid.find_gridcraft())


class Recalculo(FakeCraft):
    def recalc(self):
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=self.fake), clear=True):
            return render_grid.recalc(self.xlsx, [])

    def test_valores_y_errores_en_el_formato_de_recalc_check(self):
        res = self.recalc()
        self.assertEqual(res["motor"], "gridcraft")
        self.assertEqual(res["formulas"], 7)
        vals = {v["celda"]: v for v in res["valores"]}
        self.assertEqual(vals["C1"]["formula"], "=A1/B1")
        self.assertEqual(vals["C1"]["valor"], "#DIV/0!")
        self.assertEqual(vals["C2"]["valor"], 0.3333333333333333)  # precisión completa del JSON
        self.assertEqual(vals["C3"]["valor"], 5)
        self.assertIsInstance(vals["C3"]["valor"], int)
        self.assertEqual(vals["C4"]["valor"], "x")
        self.assertIs(vals["C5"]["valor"], True)
        self.assertEqual(vals["C6"]["valor"], "#N/A")
        self.assertEqual(vals["C7"]["valor"], "#NAME?")  # FOO no está en el catálogo: error real
        self.assertEqual([e["celda"] for e in res["errores"]], ["C1", "C6", "C7"])

    def test_fallo_del_programa_devuelve_none_con_aviso(self):
        self.instalar(self.bin, fallo=True)
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=self.fake), clear=True):
            avisos = []
            self.assertIsNone(render_grid.recalc(self.xlsx, avisos))
        self.assertTrue(avisos)

    def test_spill_devuelve_none(self):
        fila = {"0": {"0": celda("Error", "Spill", "A1:A3")}}
        self.instalar(self.bin, valores={"sheets": [{"name": "Datos", "cells": {"bands": {"0": {"rows": fila}}}}]})
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=self.fake), clear=True):
            self.assertIsNone(render_grid.recalc(self.xlsx, []))

    def test_funcion_del_catalogo_sin_calcular_devuelve_none(self):
        # #NAME? con una función que GridCraft sí conoce (FILTER) puede ser una que no implementa: no se da por bueno
        valores = libro_json({(0, 2): celda("Error", "Name", "FILTER(A1:A3,A1:A3>1)")})
        self.instalar(self.bin, valores=valores)
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=self.fake), clear=True):
            self.assertIsNone(render_grid.recalc(self.xlsx, []))

    def test_funcion_moderna_sin_prefijo_xlfn_devuelve_none(self):
        # Excel mostraría #NAME? (openpyxl escribe TEXTJOIN sin _xlfn.); GridCraft la calcularía igual
        valores = libro_json({(0, 2): celda("Text", "a,b", "TEXTJOIN(\",\",TRUE,A1:A2)")})
        self.instalar(self.bin, valores=valores)
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=self.fake), clear=True):
            self.assertIsNone(render_grid.recalc(self.xlsx, []))

    def test_tipo_desconocido_devuelve_none(self):
        valores = libro_json({(0, 2): celda("Raro", 1, "A1")})
        self.instalar(self.bin, valores=valores)
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=self.fake), clear=True):
            self.assertIsNone(render_grid.recalc(self.xlsx, []))

    def test_no_escribe_junto_al_libro_ni_en_la_carpeta_original(self):
        antes = sorted(os.listdir(self.d))
        self.recalc()
        self.assertEqual(sorted(os.listdir(self.d)), antes)


class Guardas(unittest.TestCase):
    def test_sin_xlfn(self):
        self.assertEqual(render_grid._sin_xlfn("=TEXTJOIN(\",\",TRUE,A1:A2)"), "TEXTJOIN")
        self.assertEqual(render_grid._sin_xlfn("=STDEV.S(A1:A3)"), "STDEV.S")
        self.assertIsNone(render_grid._sin_xlfn("=_XLFN.TEXTJOIN(\",\",TRUE,A1:A2)"))
        self.assertIsNone(render_grid._sin_xlfn("=_xlfn.STDEV.S(A1:A3)"))
        self.assertIsNone(render_grid._sin_xlfn("=SUM(A1:A3)+IF(A1>1,2,3)"))

    def test_coordenadas(self):
        self.assertEqual(render_grid._celda(0, 0), "A1")
        self.assertEqual(render_grid._celda(25, 0), "Z1")
        self.assertEqual(render_grid._celda(26, 0), "AA1")
        self.assertEqual(render_grid._celda(27, 9), "AB10")


class Renderizado(FakeCraft):
    def test_render_xlsx_genera_png_por_pagina(self):
        salida = os.path.join(self.d, "png")
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=self.fake), clear=True):
            hecho = render_grid.render_xlsx(self.xlsx, salida, 20, 72, [])
        self.assertIsNotNone(hecho)
        pngs, total = hecho
        self.assertEqual(total, 2)
        self.assertEqual([os.path.basename(p) for p in pngs], ["libro-p01.png", "libro-p02.png"])
        self.assertTrue(all(os.path.isfile(p) for p in pngs))

    def test_respeta_max_pages_y_avisa(self):
        avisos = []
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=self.fake), clear=True):
            pngs, total = render_grid.render_xlsx(self.xlsx, os.path.join(self.d, "p"), 1, 50, avisos)
        self.assertEqual((len(pngs), total), (1, 2))
        self.assertTrue(any("solo se rasterizaron" in a for a in avisos))

    def test_fallo_devuelve_none(self):
        self.instalar(self.bin, fallo=True)
        avisos = []
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=self.fake), clear=True):
            self.assertIsNone(render_grid.render_xlsx(self.xlsx, os.path.join(self.d, "p"), 5, 50, avisos))
        self.assertTrue(avisos)

    def test_render_py_usa_gridcraft_primero(self):
        cp = subprocess.run([sys.executable, "-I", RENDER, self.xlsx, "--outdir", os.path.join(self.d, "o")],
                            capture_output=True, text=True, env=entorno(LAZUNOX_GRIDCRAFT=self.fake))
        res = json.loads(cp.stdout)
        self.assertEqual(cp.returncode, 0, res)
        self.assertEqual((res["ok"], res["motor"], res["paginas"]), (True, "gridcraft", 2))

    def test_render_py_sin_resultado_de_gridcraft_no_hay_render(self):
        self.instalar(self.bin, fallo=True)
        cp = subprocess.run([sys.executable, "-I", RENDER, self.xlsx, "--outdir", os.path.join(self.d, "o")],
                            capture_output=True, text=True, env=entorno(LAZUNOX_GRIDCRAFT=self.fake))
        res = json.loads(cp.stdout)
        self.assertIsNone(res.get("motor"))  # GridCraft falla: no hay render, pero no revienta
        self.assertIn("GridCraft", " ".join(res["avisos"]))


class RecalcCheck(FakeCraft):
    def test_recalc_check_usa_gridcraft_y_lo_dice(self):
        valores = libro_json({(0, 2): celda("Error", "Div0", "A1/B1")})
        self.instalar(self.bin, valores=valores)
        cp = subprocess.run([sys.executable, RECALC, self.xlsx, "--json"], capture_output=True, text=True,
                            env=entorno(LAZUNOX_GRIDCRAFT=self.fake))
        res = json.loads(cp.stdout)
        self.assertEqual(res["motor"], "gridcraft")
        self.assertEqual(cp.returncode, 1)
        self.assertTrue(any(e["celda"] == "C1" and e["valor"] == "#DIV/0!" for e in res["errores"]))


@unittest.skipUnless(REAL and os.path.isfile(REAL), "define GRIDCRAFT_CLI=/ruta/a/gridcraft-cli para la prueba real")
class Real(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp(prefix="gridcraft-real-")
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Datos"
        ws["A1"], ws["A2"], ws["A3"] = 10, 20, 30
        ws["B1"] = "=SUM(A1:A3)"
        ws["B2"] = "=AVERAGE(A1:A3)"
        ws["B3"] = "=1/3"
        ws["B4"] = "=A1/0"
        ws["B5"] = "=VLOOKUP(\"zz\",A1:A3,1,FALSE)"
        self.xlsx = os.path.join(self.d, "real.xlsx")
        wb.save(self.xlsx)

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    def test_valores_reales(self):
        sha = lambda p: hashlib.sha256(open(p, "rb").read()).hexdigest()  # noqa: E731
        antes = sha(self.xlsx)
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=REAL), clear=True):
            res = render_grid.recalc(self.xlsx, [])
        self.assertIsNotNone(res)
        vals = {v["celda"]: v["valor"] for v in res["valores"]}
        self.assertEqual(vals["B1"], 60)
        self.assertAlmostEqual(vals["B2"], 20)
        self.assertAlmostEqual(vals["B3"], 1 / 3, places=12)
        self.assertEqual(vals["B4"], "#DIV/0!")
        self.assertEqual(vals["B5"], "#N/A")
        self.assertEqual(sha(self.xlsx), antes)  # el libro original no cambia

    def test_render_real(self):
        with mock.patch.dict(os.environ, entorno(LAZUNOX_GRIDCRAFT=REAL), clear=True):
            hecho = render_grid.render_xlsx(self.xlsx, os.path.join(self.d, "png"), 5, 60, [])
        self.assertIsNotNone(hecho)
        self.assertTrue(all(os.path.isfile(p) for p in hecho[0]))


if __name__ == "__main__":
    unittest.main()

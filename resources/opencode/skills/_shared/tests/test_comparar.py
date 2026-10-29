"""Pruebas de comparar_estructura.py para Word y PDF (Excel: skills/xlsx/tests/test_translate.py). Ejecutar: python -m unittest tests/test_comparar.py"""
import os
import sys
import tempfile
import unittest

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(AQUI, ".."))
import comparar_estructura as ce  # noqa: E402

try:
    import docx
except ImportError:  # pragma: no cover
    docx = None
try:
    import fitz
except ImportError:  # pragma: no cover
    fitz = None


@unittest.skipIf(docx is None, "python-docx no instalado")
class TestWord(unittest.TestCase):
    def doc(self, ruta, parrafos, filas=2, negrita=True):
        d = docx.Document()
        d.add_heading(parrafos[0], 1)
        for t in parrafos[1:]:
            p = d.add_paragraph()
            r = p.add_run(t)
            r.bold = negrita
        tabla = d.add_table(rows=filas, cols=2)
        tabla.cell(0, 0).text = "a"
        d.save(ruta)

    def test_mismo_documento_con_otro_texto_es_identico(self):
        d = tempfile.mkdtemp()
        a, b = os.path.join(d, "a.docx"), os.path.join(d, "b.docx")
        self.doc(a, ["Título", "Hola", "Mundo"])
        self.doc(b, ["Title", "Hello", "World"])
        self.assertEqual(ce.comparar(a, b)[0], [])

    def test_detecta_parrafo_de_menos_tabla_distinta_y_formato(self):
        d = tempfile.mkdtemp()
        a = os.path.join(d, "a.docx")
        self.doc(a, ["Título", "Hola", "Mundo"])
        menos, filas = os.path.join(d, "menos.docx"), os.path.join(d, "filas.docx")
        self.doc(menos, ["Title", "Hello"])
        self.doc(filas, ["Title", "Hello", "World"], filas=3)
        self.assertTrue(ce.comparar(a, menos)[0])
        self.assertTrue(ce.comparar(a, filas)[0])
        # solo formato (negrita perdida): no es diferencia de estructura, es aviso
        sin = os.path.join(d, "sin.docx")
        self.doc(sin, ["Title", "Hello", "World"], negrita=False)
        dif, avisos, _ = ce.comparar(a, sin)
        self.assertEqual(dif, [])
        self.assertTrue(avisos)


@unittest.skipIf(fitz is None, "PyMuPDF no instalado")
class TestPdf(unittest.TestCase):
    def pdf(self, ruta, paginas, texto="Hola"):
        d = fitz.open()
        for i in range(paginas):
            p = d.new_page()
            p.insert_text((72, 72), "%s %d" % (texto, i))
            p.draw_rect(fitz.Rect(100, 100, 200, 150))
        d.save(ruta)

    def test_identico_y_cambios(self):
        d = tempfile.mkdtemp()
        a, b, c = (os.path.join(d, n) for n in ("a.pdf", "b.pdf", "c.pdf"))
        self.pdf(a, 2, "Hola")
        self.pdf(b, 2, "Hello")
        self.pdf(c, 3, "Hola")
        self.assertEqual(ce.comparar(a, b)[0], [])
        self.assertTrue(ce.comparar(a, c)[0])


class TestTipos(unittest.TestCase):
    def test_tipos_distintos_o_desconocidos_salen_con_2(self):
        with self.assertRaises(SystemExit) as c:
            ce.comparar("a.docx", "b.pdf")
        self.assertEqual(c.exception.code, 2)
        with self.assertRaises(SystemExit) as c:
            ce.comparar("a.txt", "b.txt")
        self.assertEqual(c.exception.code, 2)


if __name__ == "__main__":
    unittest.main()
